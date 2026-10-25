import csv
import hashlib
import io
import json
from datetime import date, datetime
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import EmailStr, TypeAdapter
from pydantic import ValidationError as PydanticError
from sqlalchemy import select, update

from app.errors import ConflictError, ValidationError
from app.models import Employee, Team
from app.schemas import EmployeeWrite
from app.services.common import commit, record_activity
from app.services.organization import apply_reporting_changes
from app.services.personnel import employee_values

MAX_ROWS = 10000
MAX_FILE_SIZE = 5 * 1024 * 1024
HEADERS = ["email", "team_code", "department", "first_name", "last_name", "position", "hired_at"]
ALIASES = {
    "почта": "email",
    "код команды": "team_code",
    "отдел": "department",
    "имя": "first_name",
    "фамилия": "last_name",
    "должность": "position",
    "дата приёма": "hired_at",
}


def file_rows(content: bytes, filename: str):
    if not content or len(content) > MAX_FILE_SIZE:
        raise ValidationError("Выберите непустой CSV или XLSX размером до 5 МБ")
    if filename.lower().endswith(".xlsx"):
        try:
            with ZipFile(io.BytesIO(content)) as archive:
                if (
                    sum(part.file_size for part in archive.infolist()) > 30 * 1024 * 1024
                    or len(archive.infolist()) > 5000
                ):
                    raise ValidationError("Слишком большой распакованный файл XLSX")
            book = load_workbook(
                io.BytesIO(content), read_only=True, data_only=False, keep_links=False
            )
            try:
                sheet = book.worksheets[0]
                if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
                    raise ValidationError("За один импорт можно загрузить до 10 000 строк")
                rows = []
                for cells in sheet.iter_rows(max_col=20):
                    values = []
                    for cell in cells:
                        if cell.data_type == "f":
                            raise ValidationError(
                                "Замените формулы в XLSX на значения перед импортом"
                            )
                        value = cell.value
                        values.append(
                            value.date().isoformat()
                            if isinstance(value, datetime)
                            else value.isoformat()
                            if isinstance(value, date)
                            else str(value or "").strip()
                        )
                    rows.append(values)
                    if len(rows) > MAX_ROWS + 1:
                        raise ValidationError("За один импорт можно загрузить до 10 000 строк")
                return normalize_rows(rows)
            finally:
                book.close()
        except ValidationError:
            raise
        except (
            BadZipFile,
            KeyError,
            IndexError,
            ValueError,
            OSError,
            ParseError,
            InvalidFileException,
        ) as exc:
            raise ValidationError(
                "Не удалось прочитать XLSX. Проверьте файл и первый лист"
            ) from exc
    if not filename.lower().endswith(".csv"):
        raise ValidationError("Поддерживаются файлы CSV и XLSX")
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = content.decode("cp1251")
        except UnicodeDecodeError as exc:
            raise ValidationError("Сохраните CSV в кодировке UTF-8 или Windows-1251") from exc
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
        rows = []
        for row in csv.reader(io.StringIO(text), dialect, strict=True):
            rows.append(row)
            if len(rows) > MAX_ROWS + 1:
                raise ValidationError("За один импорт можно загрузить до 10 000 строк")
        return normalize_rows(rows)
    except csv.Error as exc:
        raise ValidationError("Не удалось прочитать CSV. Используйте разделитель ; или ,") from exc


def normalize_rows(rows):
    if not rows:
        raise ValidationError("Файл пуст")
    headers = [ALIASES.get(value.strip().casefold(), value.strip().casefold()) for value in rows[0]]
    nonempty = [header for header in headers if header]
    if len(nonempty) != len(set(nonempty)):
        raise ValidationError("В файле повторяются названия столбцов")
    if not {"email", "team_code"} <= set(headers):
        raise ValidationError("Нужны столбцы email и team_code. Скачайте шаблон импорта")
    result = []
    for line, values in enumerate(rows[1:], 2):
        if any(value.strip() for value in values):
            result.append(
                (
                    line,
                    {
                        header: values[index].strip() if index < len(values) else ""
                        for index, header in enumerate(headers)
                        if header
                    },
                )
            )
    if not result:
        raise ValidationError("В файле нет строк сотрудников")
    return result


def plan_import(session, content, filename):
    source = file_rows(content, filename)
    teams = {team.code: team for team in session.scalars(select(Team)).all()}
    people = {
        person.email.casefold(): person
        for person in session.scalars(
            select(Employee).where(
                Employee.email.in_([row.get("email", "").casefold() for _, row in source])
            )
        ).all()
    }
    plan, errors, preview, seen = [], [], [], set()
    fingerprint = []
    email_type = TypeAdapter(EmailStr)
    for line, row in source:
        try:
            email = str(email_type.validate_python(row.get("email", ""))).casefold()
            if email in seen:
                raise ValidationError("Email повторяется в этом файле")
            seen.add(email)
            code = row.get("team_code", "").upper()
            team = teams.get(code)
            if team is None:
                raise ValidationError(
                    f"Команда {code or 'без кода'} не найдена. Сначала создайте её"
                )
            department = row.get("department", "")
            if department and department.casefold() not in {
                str(team.department_id),
                team.department.name.casefold(),
            }:
                raise ValidationError("Отдел в строке не совпадает с отделом команды")
            person = people.get(email)
            values = None
            if person is not None:
                if person.department_id != team.department_id:
                    raise ValidationError(
                        "Сотрудник находится в другом отделе. Сначала выполните перевод"
                    )
                if person.id == team.manager_id:
                    raise ValidationError("Руководитель не может быть участником своей команды")
                name = f"{person.first_name} {person.last_name}"
                action = "unchanged" if person.team_id == team.id else "assign"
            else:
                values = EmployeeWrite(
                    first_name=row.get("first_name", ""),
                    last_name=row.get("last_name", ""),
                    email=email,
                    position=row.get("position", ""),
                    hired_at=row.get("hired_at", ""),
                    department_id=team.department_id,
                )
                name = f"{values.first_name} {values.last_name}"
                action = "create"
            plan.append((line, email, team, person, values))
            preview.append(
                {
                    "line": line,
                    "email": email,
                    "name": name,
                    "team_code": team.code,
                    "action": action,
                }
            )
            fingerprint.append(
                [
                    team.id,
                    team.version,
                    person.id if person else None,
                    person.department_id if person else None,
                    person.team_id if person else None,
                    person.manager_id if person else None,
                ]
            )
        except (ValidationError, PydanticError) as exc:
            message = (
                str(exc)
                if isinstance(exc, ValidationError)
                else "Проверьте email; для нового сотрудника обязательны first_name, last_name, "
                "position и hired_at (ГГГГ-ММ-ДД)"
            )
            errors.append({"line": line, "message": message})
    # Preflight the complete proposed graph so cycle errors are shown before confirmation.
    parents = dict(session.execute(select(Employee.id, Employee.manager_id)).all())
    for _, _, team, person, _ in plan:
        if person:
            parents[person.id] = team.manager_id or team.supervisor_id
    for line, _, _, person, _ in plan:
        if not person:
            continue
        visited, cursor = set(), person.id
        while cursor is not None:
            if cursor in visited:
                errors.append({"line": line, "message": "Назначение создаст цикл подчинённости"})
                break
            visited.add(cursor)
            cursor = parents.get(cursor)
    digest = hashlib.sha256(content + json.dumps(fingerprint, sort_keys=True).encode()).hexdigest()
    return plan, {
        "total": len(source),
        "create_count": sum(row["action"] == "create" for row in preview),
        "assign_count": sum(row["action"] == "assign" for row in preview),
        "unchanged_count": sum(row["action"] == "unchanged" for row in preview),
        "rows": preview[:100],
        "errors": errors[:100],
        "error_count": len(errors),
        "fingerprint": digest,
    }


def import_people(session, content, filename, fingerprint):
    # Serialize imports and hierarchy changes before repeating preview validation.
    session.execute(update(Team).values(version=Team.version))
    plan, preview = plan_import(session, content, filename)
    if preview["errors"]:
        raise ValidationError("Импорт отменён. Исправьте ошибки и снова проверьте файл")
    if preview["fingerprint"] != fingerprint:
        raise ConflictError("Данные изменились после проверки. Проверьте файл ещё раз")
    changes, memberships, touched = {}, {}, set()
    for _, _, team, person, values in plan:
        if person is None:
            person = Employee(**employee_values(values))
            session.add(person)
            session.flush()
        changes[person.id] = team.manager_id or team.supervisor_id
        memberships[person.id] = team.id
        touched.add(team.id)
    apply_reporting_changes(session, changes, memberships, touched)
    record_activity(
        session,
        "team",
        "Импортированы сотрудники команд",
        f"Создано: {preview['create_count']}, назначено: {preview['assign_count']}",
    )
    commit(session, "Импорт отменён: данные были изменены другим пользователем")
    return {
        key: preview[key] for key in ("total", "create_count", "assign_count", "unchanged_count")
    }
