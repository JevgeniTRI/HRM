"""Bounded CSV and multi-sheet Excel transport for the common import contract."""

import csv
import io
from datetime import date, datetime
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError as PydanticError

from app.errors import ValidationError
from app.import_schemas import (
    DepartmentImport,
    EmployeeImport,
    ImportBatch,
    SkillImport,
    TeamImport,
)
from app.services.data_import import ImportRejected

LIMIT = 5 * 1024 * 1024
RECORDS = {
    "departments": DepartmentImport,
    "employees": EmployeeImport,
    "teams": TeamImport,
    "skills": SkillImport,
}
HEADERS = list(
    dict.fromkeys(
        ["entity", *[field for model in RECORDS.values() for field in model.model_fields]]
    )
)


def template(kind):
    if kind == "csv":
        output = io.StringIO(newline="")
        csv.writer(output, delimiter=";").writerow(HEADERS)
        return output.getvalue().encode("utf-8-sig")
    book = Workbook()
    book.remove(book.active)
    for entity, model in RECORDS.items():
        sheet = book.create_sheet(entity)
        sheet.append(list(model.model_fields))
        sheet.freeze_panes = "A2"
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 28
            col[0].number_format = "@"
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def _tables(content, filename):
    if not content or len(content) > LIMIT:
        raise ValidationError("Выберите непустой CSV или XLSX размером до 5 МБ")
    if filename.lower().endswith(".csv"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = content.decode("cp1251")
            except UnicodeDecodeError as exc:
                raise ValidationError("Сохраните CSV в кодировке UTF-8 или Windows-1251") from exc
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
            rows = []
            for row in csv.reader(io.StringIO(text), dialect, strict=True):
                rows.append(row)
                if len(rows) > 10001:
                    raise ValidationError("За один импорт можно загрузить до 10 000 строк")
            return [("CSV", rows)]
        except csv.Error as exc:
            raise ValidationError(
                "Не удалось прочитать CSV. Используйте разделитель ; или ,"
            ) from exc
    if not filename.lower().endswith(".xlsx"):
        raise ValidationError("Поддерживаются CSV и XLSX")
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            if (
                len(archive.infolist()) > 5000
                or sum(part.file_size for part in archive.infolist()) > 30 * 1024 * 1024
            ):
                raise ValidationError("Слишком большой распакованный XLSX")
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            tables, count = [], 0
            for sheet in book.worksheets:
                if sheet.max_row and sheet.max_row > 10001:
                    raise ValidationError("Слишком много строк в XLSX")
                if sheet.max_column and sheet.max_column > 40:
                    raise ValidationError("В XLSX допускается до 40 столбцов на лист")
                rows = []
                for cells in sheet.iter_rows(max_col=40):
                    row = []
                    for cell in cells:
                        if cell.data_type == "f":
                            raise ValidationError("Замените формулы в XLSX на значения")
                        value = cell.value
                        row.append(
                            value.date().isoformat()
                            if isinstance(value, datetime)
                            else value.isoformat()
                            if isinstance(value, date)
                            else str(value).strip()
                            if value is not None
                            else ""
                        )
                    rows.append(row)
                    count += 1
                    if count > 10000 + len(book.worksheets):
                        raise ValidationError("За один импорт можно загрузить до 10 000 строк")
                tables.append((sheet.title, rows))
            return tables
        finally:
            book.close()
    except (
        BadZipFile,
        KeyError,
        IndexError,
        ValueError,
        OSError,
        ParseError,
        InvalidFileException,
    ) as exc:
        raise ValidationError("Не удалось прочитать XLSX. Проверьте файл") from exc


def parse(content, filename):
    entities = {name: [] for name in RECORDS}
    errors, seen, locations = [], set(), {}
    for sheet, rows in _tables(content, filename):
        if not rows:
            continue
        headers = [value.strip().lower() for value in rows[0]]
        nonempty = [value for value in headers if value]
        if len(nonempty) != len(set(nonempty)) or "external_id" not in headers:
            raise ValidationError(f"Лист {sheet}: нужны уникальные заголовки и столбец external_id")
        fixed_entity = sheet.lower() if sheet.lower() in RECORDS else None
        if not fixed_entity and "entity" not in headers:
            raise ValidationError(
                f"Лист {sheet}: нужен столбец entity либо имя листа "
                "departments, employees, teams или skills"
            )
        for line, values in enumerate(rows[1:], 2):
            if not any(value.strip() for value in values):
                continue
            raw = {
                header: values[index].strip()
                for index, header in enumerate(headers)
                if header and index < len(values) and values[index].strip()
            }
            entity = raw.pop("entity", fixed_entity or "").lower()
            location = {
                "sheet": sheet,
                "line": line,
                "entity": entity,
                "external_id": raw.get("external_id", ""),
            }
            try:
                if entity not in RECORDS or (fixed_entity and fixed_entity != entity):
                    raise ValueError(
                        "Неверный тип entity. Используйте departments, employees, teams или skills"
                    )
                if (entity, raw.get("external_id")) in seen:
                    raise ValueError("Повторяется external_id для этого типа записей")
                record = RECORDS[entity].model_validate(
                    {key: None if value == "#NULL" else value for key, value in raw.items()}
                )
                seen.add((entity, record.external_id))
                entities[entity].append(record)
                locations[(entity, record.external_id)] = location
            except PydanticError as exc:
                errors.append(
                    {
                        **location,
                        "message": "; ".join(
                            f"{'.'.join(map(str, item['loc']))}: {item['msg']}"
                            for item in exc.errors()
                        ),
                    }
                )
            except ValueError as exc:
                errors.append({**location, "message": str(exc)})
    if errors:
        raise ImportRejected(errors[:100])
    if not seen:
        raise ValidationError("В файле нет записей. Заполните строки под заголовками шаблона")
    return ImportBatch(**entities), locations


EXAMPLE = {
    "request_id": "sap-2026-10-24-001",
    "departments": [{"external_id": "D100", "name": "Разработка"}],
    "employees": [
        {
            "external_id": "E100",
            "first_name": "Анна",
            "last_name": "Иванова",
            "email": "anna.sap@example.com",
            "position": "Начальник отдела",
            "department_external_id": "D100",
            "hired_at": "2026-01-01",
        },
        {
            "external_id": "E101",
            "first_name": "Иван",
            "last_name": "Петров",
            "email": "ivan.sap@example.com",
            "position": "Руководитель команды",
            "department_external_id": "D100",
            "manager_external_id": "E100",
            "hired_at": "2026-01-01",
        },
        {
            "external_id": "E102",
            "first_name": "Мария",
            "last_name": "Смирнова",
            "email": "maria.sap@example.com",
            "position": "Разработчик",
            "department_external_id": "D100",
            "team_external_id": "T100",
            "hired_at": "2026-01-01",
        },
    ],
    "teams": [
        {
            "external_id": "T100",
            "name": "Разработка продукта",
            "code": "DEV-100",
            "department_external_id": "D100",
            "manager_external_id": "E101",
            "supervisor_external_id": "E100",
        }
    ],
    "skills": [
        {
            "external_id": "S100",
            "name": "Python",
            "category": "Технические",
            "description": "Разработка сервисов",
        }
    ],
}
